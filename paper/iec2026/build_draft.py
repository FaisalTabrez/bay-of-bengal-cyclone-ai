"""Build the IEC working draft without training or decoding final-test data."""
from pathlib import Path
import csv
import hashlib
import json
import re
import zipfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
C = json.loads((HERE/'manuscript_content.json').read_text(encoding='utf-8'))
OUT = HERE/'drafts'
TEX = HERE/'springer'
EVIDENCE = HERE/'working_evidence'
for p in [OUT, TEX/'figures', EVIDENCE]:
    p.mkdir(parents=True, exist_ok=True)

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

archives = [
    ('deterministic', 'deterministic_v0_3_results (1).zip', 'reports/colab_results_audit.json'),
    ('wind_followup', 'wind_mae_followup_v1_results (1).zip', 'reports/wind_research_export_verification.json'),
    ('uncertainty', 'uncertainty_v1_results (1).zip', 'reports/uncertainty_research_export_verification.json'),
    ('calibrated_uncertainty', 'uncertainty_v1_results (3).zip', 'reports/calibrated_export_verification.json')
]
receipts = []
for label, name, report in archives:
    path = Path.home()/'Downloads'/name
    audit = json.loads((ROOT/report).read_text())
    expected_hash = audit.get('archive_sha256')
    assert sha(path) == expected_hash, (label, 'archive hash changed')
    with zipfile.ZipFile(path) as z:
        manifest = json.loads(z.read('RUN_FILE_MANIFEST.json'))
        for filename, expected in manifest.items():
            payload = z.read(filename)
            assert len(payload) == expected['bytes'], filename
            assert hashlib.sha256(payload).hexdigest() == expected['sha256'], filename
        receipts.append({'role':label, 'archive':name, 'sha256':expected_hash, 'manifest_entries_verified':len(manifest)})
    # Development evidence is deliberately kept outside the manuscript.
    (EVIDENCE/Path(report).name).write_text(json.dumps(audit, indent=2), encoding='utf-8')

dataset = ROOT/'data/packages/bob_cyclone_training_v2.zip'
assert sha(dataset) == '3df324f20f373a8c868bc8f4b29449fb79226725e9e3705efbe898f83e67cd42'
with zipfile.ZipFile(dataset) as z:
    manifest = json.loads(z.read('training_manifest_v2.json'))
    (EVIDENCE/'dataset_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
# No Parquet file, including bob_test, is opened by this builder.
(EVIDENCE/'verification.json').write_text(json.dumps({
    'checked_on':'2026-10-30', 'archives':receipts, 'dataset_sha256':sha(dataset),
    'training_performed':False, 'test_parquet_decoded':False,
    'model_results_in_manuscript':False,
    'reason':'Author requested model results, outstanding tests and final evaluation remain on hold.'
}, indent=2), encoding='utf-8')

with zipfile.ZipFile(HERE/'sources/springer_nature_template.zip') as z:
    for member in ['sn-jnl.cls','bst/sn-mathphys-num.bst','user-manual.pdf']:
        (TEX/Path(member).name).write_bytes(z.read('sn-article-template/'+member))

rows = [
    ['Bay training (1982-2010)', '4560 / 235', '2334', '2267'],
    ['NIO incl. Bay (1982-2010)', '5557 / 289', '2848', '2742'],
    ['Selection (2011-2014)', '493 / 25', '349', '347'],
    ['Calibration (2015-2017)', '361 / 25', '298', '303'],
    ['Test metadata (2018-2025)', '1018 / 65', 'held', 'held']
]
headers = ['Role and years', 'Pairs / storms', 'Wind', 'Pressure']
caption = 'Dataset roles and observed wind/pressure target counts. NIO includes Bay training. Final-test pair/storm counts are metadata; label auditing remains on hold. Persistence-eligible counts are in Supplement S1.'
figure_caption = 'Residual prediction and frozen-mean uncertainty workflow. Reserved calibration adjusts spread; the final-test stage remains on hold.'

plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':10, 'pdf.fonttype':42})
fig, ax = plt.subplots(figsize=(7.1,3.1))
ax.set_xlim(0,10); ax.set_ylim(0,4.1); ax.axis('off')
boxes = [
    (0.12,2.45,2.65,1.30,'Historical predictors\n39 values and 39 masks\nTraining transforms'),
    (3.55,2.45,2.8,1.30,'Residual models\nSnapshot, GRU, trees\nSelection storms'),
    (7.08,2.45,2.80,1.30,'Fixed point ensemble\nThree declared seeds\nSpherical decoding'),
    (7.08,0.23,2.80,1.30,'Gaussian spread model\nFixed forecast means\nSelection checkpoint'),
    (3.55,0.23,2.80,1.30,'Reserved calibration\nThree scale factors\nUnchanged means'),
    (0.12,0.23,2.65,1.30,'Frozen final evaluation\nResults on hold\nNo performance claim')
]
for x,y,w,h,text in boxes:
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.07',
                               facecolor='#f3f3f3', edgecolor='#333333', linewidth=.9))
    ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=8.5,linespacing=1.65)
for a,b in [((2.85,3.1),(3.47,3.1)),((6.43,3.1),(7.00,3.1)),((8.48,2.36),(8.48,1.61)),
            ((7.0,.88),(6.43,.88)),((3.47,.88),(2.85,.88))]:
    ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=12,color='#333333',linewidth=1))
fig.savefig(TEX/'figures/workflow.pdf',bbox_inches='tight',pad_inches=.07)
fig.savefig(TEX/'figures/workflow.png',dpi=300,bbox_inches='tight',pad_inches=.07)
plt.close(fig)

for eq in C.get('equations',[]):
    fig=plt.figure(figsize=(7,.55))
    fig.text(.5,.5,'$'+eq['latex']+r'\quad ('+eq['label']+')$',ha='center',va='center',fontsize=13)
    fig.savefig(TEX/('figures/equation_'+eq['label']+'.png'),dpi=300,bbox_inches='tight',pad_inches=.08)
    plt.close(fig)

refnum = {r['key']:i+1 for i,r in enumerate(C['references'])}
def cite_text(s):
    return re.sub(r'\[@([^\]]+)\]', lambda m:'['+', '.join(str(refnum[k.lstrip('@')]) for k in m.group(1).split(';'))+']',s)
def esc(s):
    for a,b in [('\\',r'\textbackslash{}'),('&',r'\&'),('%',r'\%'),('$',r'\$'),('#',r'\#'),('_',r'\_')]:
        s=s.replace(a,b)
    return s
def tex_text(s):
    parts = re.split(r'(\[@[^\]]+\])',s)
    return ''.join(r'\cite{'+','.join(k.lstrip('@') for k in p[2:-1].split(';'))+'}' if p.startswith('[@') else esc(p) for p in parts)

tex = r'''% Official Springer Nature journal class; IEC author guidelines take precedence.
% Results and outstanding tests remain on hold at the authors' request.
\documentclass[pdflatex,sn-mathphys-num,referee,a4paper]{sn-jnl}
\usepackage{graphicx}
\usepackage{booktabs}
\usepackage{amsmath}
\usepackage{xcolor}
\begin{document}
'''
tex += '\\title{'+esc(C['title'])+'}\n'
tex += r'\author*[1]{\fnm{Md. Faisal} \sur{Tabrez}}\email{'+C['corresponding_email']+'}\n'
tex += r'\author[1]{\fnm{K.} \sur{Akhilesh}}\email{'+C['second_author_email']+'}\n'
tex += r'\author[1]{\fnm{T. V.} \sur{Rajinikanth}}'+'\n'
tex += r'\affil[1]{\orgdiv{Department of Computer Science and Engineering}, \orgname{Mahatma Gandhi Institute of Technology}, \orgaddress{\city{Hyderabad}, \state{Telangana}, \country{India}}}'+'\n'
tex += '\\abstract{'+esc(C['abstract'])+'}\n\\keywords{'+', '.join(C['keywords'])+'}\n\\maketitle\n'
tex += r'\noindent\textbf{Notation.} $t$: current observation time; NIO: North Indian Ocean; ERA5: ECMWF reanalysis; OISST: optimum interpolation sea surface temperature.'+'\n'
for i,sec in enumerate(C['sections']):
    tex += '\n\\section{'+esc(sec['title'])+'}\n'
    for j,p in enumerate(sec['paragraphs']):
        tex += tex_text(p)+'\n\n'
        for eq in C.get('equations',[]):
            if eq['section_index']==i and eq['after_paragraph']==j:
                tex += '\\begin{equation}\n'+eq['latex']+'\n\\end{equation}\n'
    if i == 1:
        tex += '\n\\begin{table}[htbp]\n\\caption{'+esc(caption)+'}\\label{tab:data}\n\\centering\n\\begin{tabular}{p{43mm}rrr}\n\\toprule\nRole and years & Pairs / storms & Wind & Pressure \\\\\n\\midrule\n'
        tex += '\n'.join(' & '.join(esc(c) for c in row)+r' \\' for row in rows)
        tex += '\n\\bottomrule\n\\end{tabular}\n\\end{table}\n'
    if i == 3:
        tex += '\n\\begin{figure}[htbp]\n\\centering\n\\includegraphics[width=\\linewidth]{figures/workflow.pdf}\n\\caption{'+esc(figure_caption)+'}\\label{fig:workflow}\n\\end{figure}\n'
tex += '\n\\backmatter\n'
for heading,para in C['declarations']:
    tex += '\n\\subsection*{'+heading+'}\n'+esc(para)+'\n'
tex += '\n\\begin{thebibliography}{'+str(len(C['references']))+'}\n\\interlinepenalty=10000\n'
for r in C['references']:
    tex += '\\bibitem{'+r['key']+'}\n'+esc(r['text'])+' '
    if 'doi' in r: tex += '\\url{https://doi.org/'+r['doi']+'}'
    else: tex += '\\url{'+r['url']+'}'
    tex += '\n\n'
tex += '\\end{thebibliography}\n\\section*{List of captions}\nTable 1. '+esc(caption)+'\n\nFigure 1. '+esc(figure_caption)+'\n\\end{document}\n'
(TEX/'main.tex').write_text(tex,encoding='utf-8')

def new_doc():
    d=Document()
    sec=d.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=Mm(23); sec.bottom_margin=Mm(23)
    sec.left_margin=Mm(25); sec.right_margin=Mm(25)
    for style in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        s=d.styles[style]; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
        fonts=s.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower():del fonts.attrib[attr]
            for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Times New Roman')
    # Remove inherited Word template rules, including the Title blue border.
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing=2
    d.styles['Normal'].paragraph_format.space_after=Pt(3)
    d.styles['Title'].font.size=Pt(16); d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.line_spacing=1.1
    d.styles['Title'].paragraph_format.space_after=Pt(8)
    for s in ['Heading 1','Heading 2']:
        d.styles[s].font.size=Pt(12)
        d.styles[s].paragraph_format.space_before=Pt(10)
        d.styles[s].paragraph_format.space_after=Pt(4)
        d.styles[s].paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(10)
    d.styles['Caption'].paragraph_format.line_spacing=2
    d.core_properties.author='; '.join(C['authors'])
    d.core_properties.title=C['title']
    d.core_properties.subject='41st Indian Engineering Congress 2026 working draft'
    # Page numbers identify manuscript pages; no decorative running header.
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    return d

def front(d):
    d.add_paragraph(C['title'],'Title')
    p=d.add_paragraph(', '.join(C['authors']))
    p.paragraph_format.line_spacing=1
    p=d.add_paragraph(C['affiliation']); p.paragraph_format.line_spacing=1
    p=d.add_paragraph('Correspondence: '+C['corresponding_email']); p.paragraph_format.line_spacing=1
    p=d.add_paragraph('K. Akhilesh: '+C['second_author_email']); p.paragraph_format.line_spacing=1

def add_cited(d,s):
    p=d.add_paragraph()
    # IEC-linked guidelines require sequential superscript reference numerals.
    for part in re.split(r'(\[@[^\]]+\])',s):
        if part.startswith('[@'):
            r=p.add_run(','.join(str(refnum[k.lstrip('@')]) for k in part[2:-1].split(';')))
            r.font.superscript=True
        else:p.add_run(part)
    return p

def table(d):
    p=d.add_paragraph('Table 1. '+caption,'Caption'); p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=[Mm(65),Mm(38),Mm(25),Mm(25)]
    for cell,w,text in zip(t.rows[0].cells,widths,headers):
        cell.width=w; cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for cell,w,text in zip(cells,widths,row):cell.width=w;cell.text=text
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        if ri==0:trPr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcPr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');tcPr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=2;p.paragraph_format.space_after=Pt(1)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)

d=new_doc();front(d)
d.add_heading('Abstract',1);d.add_paragraph(C['abstract'])
d.add_paragraph('Keywords: '+'; '.join(C['keywords']))
d.add_paragraph('Notation: t denotes the current observation time; NIO denotes the North Indian Ocean; ERA5 denotes ECMWF reanalysis; OISST denotes optimum interpolation sea surface temperature.')
for i,sec in enumerate(C['sections'],1):
    d.add_heading(str(i)+' '+sec['title'],1)
    for j,p in enumerate(sec['paragraphs']):
        add_cited(d,p)
        for eq in C.get('equations',[]):
            if eq['section_index']==i-1 and eq['after_paragraph']==j:
                paragraph=d.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(TEX/('figures/equation_'+eq['label']+'.png')),width=Mm(154 if eq['label']=='1' else 100))
    if i==2:table(d)
    if i==4:
        p=d.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(TEX/'figures/workflow.png'),width=Mm(158))
        d.add_paragraph('Figure 1. '+figure_caption,'Caption')
for h,p in C['declarations']:d.add_heading(h,2);d.add_paragraph(p)
d.add_heading('References',1)
for i,r in enumerate(C['references'],1):
    add_cited(d,str(i)+'. '+r['text']+' '+('https://doi.org/'+r['doi'] if 'doi' in r else r['url']))
d.add_heading('List of captions',1)
d.add_paragraph('Table 1. '+caption)
d.add_paragraph('Figure 1. '+figure_caption)
d.save(OUT/'IEC2026_Paper_Draft.docx')

a=new_doc();front(a)
# IEC specifies the extended abstract's word limit, not double spacing. A compact
# abstract layout keeps the initial submission readable in two pages.
a.styles['Normal'].paragraph_format.line_spacing=1.3
a.add_paragraph('Keywords: '+'; '.join(C['keywords']))
for sec in C['extended_abstract']:
    a.add_heading(sec['title'],1)
    for p in sec['paragraphs']:a.add_paragraph(p)
a.save(OUT/'IEC2026_Extended_Abstract.docx')

md=['# '+C['title'],'',', '.join(C['authors']),'',C['affiliation'],'','Correspondence: '+C['corresponding_email'],'','## Abstract','',C['abstract'],'','Keywords: '+'; '.join(C['keywords'])]
for i,sec in enumerate(C['sections'],1):
    md.extend(['','## '+str(i)+' '+sec['title'],''])
    for j,p in enumerate(sec['paragraphs']):
        md.extend([cite_text(p),''])
        for eq in C.get('equations',[]):
            if eq['section_index']==i-1 and eq['after_paragraph']==j:
                md.extend(['```math',eq['latex'],'```',''])
    if i==2:
        md.extend(['Table 1. '+caption,'','| Role and years | Pairs / storms | Wind | Pressure |','|---|---:|---:|---:|'])
        md.extend('| '+' | '.join(row)+' |' for row in rows)
    if i==4:md.extend(['','![Workflow](springer/figures/workflow.png)','',figure_caption])
for h,p in C['declarations']:md.extend(['','## '+h,'',p])
md.extend(['','## References',''])
for i,r in enumerate(C['references'],1):md.append(str(i)+'. '+r['text']+' '+('https://doi.org/'+r['doi'] if 'doi' in r else r['url']))
(HERE/'paper_draft.md').write_text('\n'.join(md)+'\n',encoding='utf-8')

words=lambda s:len(re.findall(r"\b[\w'-]+\b",s))
abstract_text=' '.join([C['title'],*C['authors'],C['affiliation'],*C['keywords']]+[s['title']+' '+' '.join(s['paragraphs']) for s in C['extended_abstract']])
paper_words=' '.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
figure_words=' '.join(box[-1] for box in boxes)
counts={'synopsis_words':words(C['abstract']),'synopsis_whitespace_words':len(C['abstract'].split()),'extended_abstract_words_conservative':words(abstract_text),'paper_body_tables_references_and_declarations_words':words(paper_words),'figure_label_words':words(figure_words),'paper_words_including_references_and_declarations':words(paper_words)+words(figure_words),'equation_objects':len(C.get('equations',[])),'count_note':'Regex word-like count of all Word body paragraphs and table cells plus figure labels; headers/footers and formula symbols excluded. Not a publisher-certified count.','figures':1,'tables':1}
assert counts['synopsis_words']==150 and counts['synopsis_whitespace_words']==150, counts
assert counts['extended_abstract_words_conservative']<=1000,counts
assert counts['paper_words_including_references_and_declarations']<=3000,counts
(HERE/'draft_checks.json').write_text(json.dumps(counts,indent=2),encoding='utf-8')
print(json.dumps({'outputs':[str(p) for p in OUT.glob('*.docx')],'counts':counts,'archives':receipts},indent=2))
